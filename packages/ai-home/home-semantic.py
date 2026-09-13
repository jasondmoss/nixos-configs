  #!/usr/bin/env python3
"""home-semantic — meaning-based search over selected home directories.

Companion to home-search-mcp.py (exact-word Recoll search). Documents under
services.ai-home.semanticPaths are split into passages, embedded with the
local Ollama server (nomic-embed-text) and stored in an sqlite-vec database
next to the Recoll index. Three entry points share this file so that queries
and documents are always embedded by the same code:

    home-semantic index [--rebuild]      incremental (re)index; run by the
                                         ai-home-semantic-index timer
    home-semantic serve                  MCP server over stdio, exposed by
                                         mcp-proxy as /servers/semantic/mcp
    home-semantic query TEXT [-n N] [-d DIR]   shell helper (`ai-semantic`)

Everything runs inside the ai-home sandbox (see ../../ai-home.nix): paths in
services.ai-home.hiddenPaths do not exist in the mount namespace, and the
only reachable network address is loopback, so nothing can leave the machine.
Configuration is the JSON file named by $AI_HOME_SEMANTIC_CONFIG (generated
by the module).
"""

from __future__ import annotations

import argparse
import fnmatch
import json
import logging
import os
import re
import sqlite3
import stat
import sys
import time
from typing import Any

import sqlite_vec

log = logging.getLogger("home-semantic")

# nomic-embed-text is trained with task prefixes; documents and queries must
# use different ones or retrieval quality drops noticeably.
DOC_PREFIX = "search_document: "
QUERY_PREFIX = "search_query: "

# Suffixes with a dedicated extractor. Anything else listed in
# semanticSuffixes is read as plain UTF-8 text.
STRUCTURED = {
    ".pdf", ".docx", ".pptx", ".xlsx", ".odt", ".ods", ".odp",
    ".epub", ".rtf", ".html", ".htm",
}

DEFAULTS = {
    "chunk_tokens": 400,   # ≈ words+punctuation; well under the model's 2048
    "chunk_overlap": 60,
    "max_file_mb": 64,
    "batch_size": 32,
    "keep_alive": "5m",
    "min_alnum": 40,       # passages with fewer letters/digits are noise
}


def _load_config() -> dict[str, Any]:
    path = os.environ.get("AI_HOME_SEMANTIC_CONFIG")
    if not path:
        sys.exit("AI_HOME_SEMANTIC_CONFIG is not set (this program is meant to "
                 "be started through the ai-home.nix wrappers)")
    with open(path, encoding="utf-8") as fh:
        cfg = json.load(fh)
    for key, value in DEFAULTS.items():
        cfg.setdefault(key, value)
    return cfg


CFG = _load_config()


# ─── storage ───────────────────────────────────────────────────────────────

def _connect(create: bool = False) -> sqlite3.Connection:
    db = CFG["db"]
    if create:
        os.makedirs(os.path.dirname(db), exist_ok=True)
    elif not os.path.exists(db):
        raise FileNotFoundError(
            "the semantic index has not been built yet; it is created by the "
            "ai-home-semantic-index timer (systemctl start ai-home-semantic-index)"
        )
    conn = sqlite3.connect(db, timeout=60)
    conn.enable_load_extension(True)
    sqlite_vec.load(conn)
    conn.enable_load_extension(False)
    conn.execute("PRAGMA busy_timeout = 60000")
    conn.execute("PRAGMA foreign_keys = ON")
    if create:
        conn.execute("PRAGMA journal_mode = WAL")
    return conn


def _init_schema(conn: sqlite3.Connection, dims: int) -> None:
    conn.executescript(
        f"""
        CREATE TABLE IF NOT EXISTS meta (
            key   TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS files (
            id         INTEGER PRIMARY KEY,
            path       TEXT    NOT NULL UNIQUE,
            mtime_ns   INTEGER NOT NULL,
            size       INTEGER NOT NULL,
            indexed_at INTEGER NOT NULL,
            chunks     INTEGER NOT NULL DEFAULT 0,
            error      TEXT
        );
        CREATE TABLE IF NOT EXISTS chunks (
            id      INTEGER PRIMARY KEY,
            file_id INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
            ordinal INTEGER NOT NULL,
            locator TEXT,
            text    TEXT    NOT NULL
        );
        CREATE INDEX IF NOT EXISTS chunks_file ON chunks(file_id);
        CREATE VIRTUAL TABLE IF NOT EXISTS chunks_vec USING vec0(
            chunk_id INTEGER PRIMARY KEY,
            embedding float[{int(dims)}] distance_metric=cosine
        );
        """
    )


def _drop_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        "DROP TABLE IF EXISTS chunks_vec; DROP TABLE IF EXISTS chunks;"
        "DROP TABLE IF EXISTS files; DROP TABLE IF EXISTS meta;"
    )


def _meta(conn: sqlite3.Connection) -> dict[str, str]:
    try:
        return dict(conn.execute("SELECT key, value FROM meta"))
    except sqlite3.OperationalError:
        return {}


def _set_meta(conn: sqlite3.Connection, **values: Any) -> None:
    conn.executemany(
        "INSERT INTO meta(key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        [(k, str(v)) for k, v in values.items()],
    )


def _delete_file_rows(conn: sqlite3.Connection, file_id: int) -> None:
    # vec0 is a virtual table: no cascade, delete its rows explicitly.
    conn.execute(
        "DELETE FROM chunks_vec WHERE chunk_id IN "
        "(SELECT id FROM chunks WHERE file_id = ?)", (file_id,))
    conn.execute("DELETE FROM chunks WHERE file_id = ?", (file_id,))


# ─── embeddings ────────────────────────────────────────────────────────────

def _client():
    import ollama
    return ollama.Client(host=CFG["ollama_url"], timeout=300)


def _embed(client, texts: list[str], prefix: str) -> list[list[float]]:
    response = client.embed(
        model=CFG["model"],
        input=[prefix + t for t in texts],
        truncate=True,
        keep_alive=CFG["keep_alive"],
    )
    vectors = list(response.embeddings)
    if len(vectors) != len(texts):
        raise RuntimeError(f"Ollama returned {len(vectors)} embeddings for {len(texts)} inputs")
    return vectors


def _backend_errors() -> tuple[type[BaseException], ...]:
    import httpx
    import ollama
    return (ollama.ResponseError, ollama.RequestError, httpx.HTTPError, ConnectionError, OSError)


# ─── file discovery ────────────────────────────────────────────────────────

def _under(path: str, prefixes: list[str]) -> bool:
    return any(path == p or path.startswith(p + "/") for p in prefixes)


def _skipped_name(name: str) -> bool:
    return any(fnmatch.fnmatchcase(name, pat) for pat in CFG.get("skipped_names", []))


def _walk() -> dict[str, tuple[int, int]]:
    """Regular files under the roots worth embedding: path -> (mtime_ns, size)."""
    hidden = [p.rstrip("/") for p in CFG.get("hidden", []) + CFG.get("unindexed", [])]
    suffixes = tuple(s.lower() for s in CFG["suffixes"])
    max_bytes = int(CFG["max_file_mb"]) * 1024 * 1024
    found: dict[str, tuple[int, int]] = {}

    for root in CFG["roots"]:
        root = root.rstrip("/")
        if not os.path.isdir(root):
            log.info("root %s does not exist, skipping", root)
            continue
        if _under(root, hidden):
            log.warning("root %s is hidden or unindexed, skipping", root)
            continue
        for dirpath, dirnames, filenames in os.walk(root, onerror=lambda e: log.debug("%s", e)):
            dirnames[:] = sorted(
                d for d in dirnames
                if not _skipped_name(d) and not _under(os.path.join(dirpath, d), hidden)
            )
            for name in filenames:
                if not name.lower().endswith(suffixes) or _skipped_name(name):
                    continue
                path = os.path.join(dirpath, name)
                try:
                    st = os.lstat(path)
                except OSError:
                    continue
                # lstat: symlinks are not regular files, so they are skipped,
                # matching followLinks=0 in the Recoll config.
                if not stat.S_ISREG(st.st_mode) or st.st_size == 0 or st.st_size > max_bytes:
                    continue
                found[path] = (st.st_mtime_ns, st.st_size)
    return found


# ─── text extraction ───────────────────────────────────────────────────────

def _html_text(markup: bytes | str) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(markup, "html.parser")
    for tag in soup(["script", "style", "noscript", "template", "svg"]):
        tag.decompose()
    return soup.get_text("\n")


def _odf_text(path: str) -> str:
    from odf import teletype
    from odf import text as odftext
    from odf.opendocument import load

    para = {odftext.P().qname, odftext.H().qname}
    parts: list[str] = []

    def visit(node) -> None:
        for child in getattr(node, "childNodes", []):
            if getattr(child, "qname", None) in para:
                parts.append(teletype.extractText(child))
            else:
                visit(child)

    visit(load(path).body)
    return "\n".join(parts)


def _extract(path: str) -> list[tuple[str | None, str]]:
    """Text parts of a file as (locator, text); locator names a page/slide/sheet."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                return []
            return [(f"page {i}", page.get_text("text")) for i, page in enumerate(doc, 1)]

    if ext == ".docx":
        import docx2txt
        return [(None, docx2txt.process(path))]

    if ext == ".pptx":
        from pptx import Presentation
        parts = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
            parts.append((f"slide {i}", "\n".join(texts)))
        return parts

    if ext == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            parts = []
            for sheet in book.worksheets:
                rows = (
                    "\t".join("" if v is None else str(v) for v in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                parts.append((f"sheet {sheet.title}", "\n".join(r for r in rows if r.strip())))
            return parts
        finally:
            book.close()

    if ext in (".odt", ".ods", ".odp"):
        return [(None, _odf_text(path))]

    if ext == ".epub":
        import ebooklib
        from ebooklib import epub
        book = epub.read_epub(path, options={"ignore_ncx": True})
        return [
            (f"section {i}", _html_text(item.get_content()))
            for i, item in enumerate(book.get_items_of_type(ebooklib.ITEM_DOCUMENT), 1)
        ]

    if ext in (".html", ".htm"):
        with open(path, "rb") as fh:
            return [(None, _html_text(fh.read()))]

    if ext == ".rtf":
        from striprtf.striprtf import rtf_to_text
        with open(path, encoding="utf-8", errors="ignore") as fh:
            return [(None, rtf_to_text(fh.read()))]

    # Plain text of any kind (md, txt, csv, ...). A NUL byte means binary.
    with open(path, "rb") as fh:
        raw = fh.read()
    if b"\0" in raw:
        return []
    return [(None, raw.decode("utf-8", errors="replace"))]


_WS = re.compile(r"[ \t\f\v ]+")
_BLANKS = re.compile(r"\n\s*\n+")


def _clean(text: str) -> str:
    text = _WS.sub(" ", text.replace("\r", ""))
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANKS.sub("\n\n", text).strip()


def _alnum(text: str) -> int:
    return sum(ch.isalnum() for ch in text)


# ─── chunking ──────────────────────────────────────────────────────────────

_TOKEN = re.compile(r"\w+|[^\w\s]", re.UNICODE)
_SENTENCE = re.compile(r"(?<=[.!?])\s+(?=\S)")


def _ntok(text: str) -> int:
    return len(_TOKEN.findall(text))


def _units(text: str, size: int) -> list[tuple[str, int, bool]]:
    """Sentences as (text, tokens, starts_paragraph), oversized ones hard-split."""
    units: list[tuple[str, int, bool]] = []
    for para in text.split("\n\n"):
        first = True
        for sentence in _SENTENCE.split(para.replace("\n", " ")):
            sentence = sentence.strip()
            if not sentence:
                continue
            words = sentence.split()
            step = max(1, size // 2)  # tables, code, URLs: no sentence marks
            for i in range(0, len(words), step):
                piece = " ".join(words[i:i + step])
                units.append((piece, _ntok(piece), first))
                first = False
    return units


def _chunk(text: str) -> list[str]:
    size, overlap = int(CFG["chunk_tokens"]), int(CFG["chunk_overlap"])
    chunks: list[str] = []
    current: list[tuple[str, int, bool]] = []
    used = 0

    def flush() -> None:
        out = ""
        for piece, _, starts_para in current:
            out += ("\n" if starts_para and out else " " if out else "") + piece
        chunks.append(out)

    for unit in _units(text, size):
        if current and used + unit[1] > size:
            flush()
            kept: list[tuple[str, int, bool]] = []
            kept_tokens = 0
            for piece in reversed(current):
                if kept_tokens + piece[1] > overlap:
                    break
                kept.insert(0, piece)
                kept_tokens += piece[1]
            current, used = kept, kept_tokens
        current.append(unit)
        used += unit[1]
    if current:
        flush()
    return chunks


# ─── indexing ──────────────────────────────────────────────────────────────

def _file_chunks(path: str) -> list[tuple[int, str | None, str]]:
    rows: list[tuple[int, str | None, str]] = []
    for locator, text in _extract(path):
        text = _clean(text)
        if _alnum(text) < int(CFG["min_alnum"]):
            continue
        for chunk in _chunk(text):
            if _alnum(chunk) >= int(CFG["min_alnum"]):
                rows.append((len(rows), locator, chunk))
    return rows


def cmd_index(rebuild: bool) -> int:
    started = time.time()
    client = _client()
    try:
        dims = len(_embed(client, ["dimension probe"], DOC_PREFIX)[0])
    except _backend_errors() as exc:
        log.error("cannot embed with %s at %s: %s", CFG["model"], CFG["ollama_url"], exc)
        log.error("is Ollama running and the model pulled? (ollama pull %s)", CFG["model"])
        return 1

    conn = _connect(create=True)
    meta = _meta(conn)
    if meta and (rebuild or meta.get("model") != CFG["model"] or meta.get("dims") != str(dims)):
        log.info("rebuilding index (model %s → %s, %s dims)", meta.get("model"), CFG["model"], dims)
        _drop_schema(conn)
        meta = {}
    _init_schema(conn, dims)
    with conn:
        _set_meta(conn, model=CFG["model"], dims=dims, run_started=int(started),
                  roots=json.dumps(CFG["roots"]))

    found = _walk()
    known = {
        path: (fid, mtime_ns, size)
        for fid, path, mtime_ns, size in conn.execute("SELECT id, path, mtime_ns, size FROM files")
    }
    removed = [fid for path, (fid, _, _) in known.items() if path not in found]
    todo = sorted(path for path, sig in found.items()
                  if path not in known or known[path][1:] != sig)
    log.info("%d candidate files: %d unchanged, %d to (re)index, %d removed",
             len(found), len(found) - len(todo), len(todo), len(removed))

    with conn:
        for fid in removed:
            _delete_file_rows(conn, fid)
            conn.execute("DELETE FROM files WHERE id = ?", (fid,))

    batch = int(CFG["batch_size"])
    done = failed = embedded = 0
    status = 0
    try:
        for n, path in enumerate(todo, 1):
            mtime_ns, size = found[path]
            error: str | None = None
            rows: list[tuple[int, str | None, str]] = []
            vectors: list[list[float]] = []
            try:
                rows = _file_chunks(path)
                for i in range(0, len(rows), batch):
                    vectors.extend(_embed(client, [r[2] for r in rows[i:i + batch]], DOC_PREFIX))
                if not rows:
                    error = "no extractable text"
            except _backend_errors() as exc:
                log.error("embedding backend failed on %s: %s — stopping, progress is kept", path, exc)
                status = 1
                break
            except Exception as exc:  # noqa: BLE001 — one bad file must not stop the run
                rows, vectors = [], []
                error = f"{type(exc).__name__}: {exc}"[:300]
                log.warning("%s: %s", path, error)

            with conn:
                old = conn.execute("SELECT id FROM files WHERE path = ?", (path,)).fetchone()
                if old:
                    _delete_file_rows(conn, old[0])
                conn.execute(
                    "INSERT INTO files(path, mtime_ns, size, indexed_at, chunks, error) "
                    "VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(path) DO UPDATE SET "
                    "mtime_ns = excluded.mtime_ns, size = excluded.size, "
                    "indexed_at = excluded.indexed_at, chunks = excluded.chunks, "
                    "error = excluded.error",
                    (path, mtime_ns, size, int(time.time()), len(rows), error),
                )
                fid = conn.execute("SELECT id FROM files WHERE path = ?", (path,)).fetchone()[0]
                for (ordinal, locator, text), vector in zip(rows, vectors):
                    cur = conn.execute(
                        "INSERT INTO chunks(file_id, ordinal, locator, text) VALUES (?, ?, ?, ?)",
                        (fid, ordinal, locator, text))
                    conn.execute("INSERT INTO chunks_vec(chunk_id, embedding) VALUES (?, ?)",
                                 (cur.lastrowid, sqlite_vec.serialize_float32(vector)))
            done += 1
            failed += error is not None
            embedded += len(rows)
            if n % 25 == 0 or n == len(todo):
                log.info("%d/%d files, %d passages embedded, %d without text (%.0fs)",
                         n, len(todo), embedded, failed, time.time() - started)
    except KeyboardInterrupt:
        log.warning("interrupted, progress is kept")
        status = 1

    with conn:
        total_files, total_chunks = conn.execute(
            "SELECT (SELECT COUNT(*) FROM files WHERE chunks > 0), (SELECT COUNT(*) FROM chunks)"
        ).fetchone()
        _set_meta(conn, run_finished=int(time.time()), run_status=status,
                  run_files=done, run_failed=failed, run_passages=embedded,
                  run_removed=len(removed), documents=total_files, passages=total_chunks)
    conn.execute("PRAGMA optimize")
    conn.close()
    log.info("done: %d documents, %d passages in the index (%.0fs)",
             total_files, total_chunks, time.time() - started)
    return status


# ─── search ────────────────────────────────────────────────────────────────

def _iso(epoch: float | int | None) -> str | None:
    try:
        return time.strftime("%Y-%m-%d %H:%M", time.localtime(int(epoch)))
    except (TypeError, ValueError, OverflowError):
        return None


def _like_prefix(path: str) -> str:
    return re.sub(r"([\\%_])", r"\\\1", path) + "/%"


def semantic_search(
    query: str,
    max_results: int = 10,
    directory: str | None = None,
) -> list[dict[str, Any]]:
    """Find documents by meaning, not by exact words.

    Describe what you are looking for in a sentence ("the contract about
    website maintenance for the ski resort", "notes on setting up a
    WireGuard tunnel") and the closest passages come back, best first, one
    entry per document: path, title, modification time, where in the file
    the passage sits (page/slide/sheet), a similarity score (1 = identical
    meaning, below ~0.45 is usually noise) and the passage itself. Covers
    the configured document folders only (see semantic_index_status); for
    exact terms, file names, code or anything outside those folders use the
    full-text search_files tool instead. `directory` restricts results to a
    subtree (absolute or relative to the home directory).
    """
    max_results = max(1, min(int(max_results), 50))
    conn = _connect()
    try:
        try:
            vector = _embed(_client(), [query.strip()], QUERY_PREFIX)[0]
        except _backend_errors() as exc:
            raise RuntimeError(
                f"cannot embed the query: Ollama at {CFG['ollama_url']} with model "
                f"{CFG['model']} failed ({exc})") from exc
        sql = ("SELECT chunk_id, distance FROM chunks_vec "
               "WHERE embedding MATCH ? AND k = ?")
        params: list[Any] = [sqlite_vec.serialize_float32(vector), max_results * 8]
        if directory:
            home = CFG.get("home", os.path.expanduser("~"))
            base = directory if directory.startswith("/") else os.path.join(home, directory)
            base = os.path.normpath(base)
            sql += (" AND chunk_id IN (SELECT c.id FROM chunks c JOIN files f ON f.id = c.file_id"
                    " WHERE f.path LIKE ? ESCAPE '\\')")
            params.append(_like_prefix(base))
        hits = conn.execute(sql, params).fetchall()

        results: list[dict[str, Any]] = []
        seen: dict[int, dict[str, Any]] = {}
        for chunk_id, distance in hits:  # already ordered by distance
            row = conn.execute(
                "SELECT f.id, f.path, f.mtime_ns, c.locator, c.text FROM chunks c "
                "JOIN files f ON f.id = c.file_id WHERE c.id = ?", (chunk_id,)).fetchone()
            if row is None:
                continue
            fid, path, mtime_ns, locator, text = row
            if fid in seen:
                seen[fid]["matching_passages"] += 1
                continue
            # Anything the sandbox hides is unreadable here, even if a stale
            # entry survived a hiddenPaths change.
            if not os.access(path, os.R_OK):
                continue
            entry = {
                "path": path,
                "title": os.path.basename(path),
                "modified": _iso(mtime_ns / 1e9),
                "location": locator,
                "score": round(1.0 - float(distance), 3),
                "passage": re.sub(r"\s+", " ", text)[:1200],
                "matching_passages": 1,
            }
            seen[fid] = entry
            results.append(entry)
            if len(results) >= max_results:
                break
        return results
    finally:
        conn.close()


def semantic_index_status() -> dict[str, Any]:
    """Report the semantic index: folders covered, document and passage counts, embedding model, last run."""
    status: dict[str, Any] = {
        "folders": CFG["roots"],
        "file_types": CFG["suffixes"],
        "model": CFG["model"],
    }
    try:
        conn = _connect()
    except FileNotFoundError as exc:
        status.update(indexed=False, message=str(exc))
        return status
    try:
        meta = _meta(conn)
        status.update(
            indexed=True,
            documents=int(meta.get("documents", 0)),
            passages=int(meta.get("passages", 0)),
            documents_without_text=conn.execute(
                "SELECT COUNT(*) FROM files WHERE chunks = 0").fetchone()[0],
            last_run_started=_iso(meta.get("run_started")),
            last_run_finished=_iso(meta.get("run_finished")),
            last_run_ok=meta.get("run_status") == "0",
            last_run_reindexed=int(meta.get("run_files", 0)),
            last_run_removed=int(meta.get("run_removed", 0)),
        )
    finally:
        conn.close()
    return status


def cmd_serve() -> int:
    from fastmcp import FastMCP

    mcp = FastMCP(
        "home-semantic",
        instructions=(
            "Meaning-based search over the user's document folders. Use "
            "semantic_search when the request describes a topic, a concept or "
            "the gist of a document rather than its exact words; use the "
            "full-text search_files tool for exact terms, file names and code. "
            "Read the returned paths with the filesystem tools."
        ),
    )
    mcp.tool(semantic_search)
    mcp.tool(semantic_index_status)
    mcp.run()
    return 0


def cmd_query(text: str, n: int, directory: str | None) -> int:
    try:
        hits = semantic_search(text, n, directory)
    except (FileNotFoundError, RuntimeError) as exc:
        print(f"ai-semantic: {exc}", file=sys.stderr)
        return 1
    for hit in hits:
        where = f"  ({hit['location']})" if hit["location"] else ""
        print(f"{hit['score']:.3f}  {hit['path']}{where}")
        print(f"       {hit['passage'][:300]}")
    return 0


def main() -> int:
    logging.basicConfig(stream=sys.stderr, level=logging.INFO,
                        format="%(levelname)s %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)  # one line per embed call otherwise
    parser = argparse.ArgumentParser(prog="home-semantic", description=__doc__.split("\n\n")[0])
    sub = parser.add_subparsers(dest="cmd", required=True)
    p_index = sub.add_parser("index", help="incrementally (re)build the index")
    p_index.add_argument("--rebuild", action="store_true", help="drop the index and start over")
    sub.add_parser("serve", help="run the MCP server on stdio")
    p_query = sub.add_parser("query", help="search from the shell")
    p_query.add_argument("text")
    p_query.add_argument("-n", type=int, default=10)
    p_query.add_argument("-d", "--directory")
    args = parser.parse_args()

    if args.cmd == "index":
        return cmd_index(args.rebuild)
    if args.cmd == "serve":
        return cmd_serve()
    return cmd_query(args.text, args.n, args.directory)


if __name__ == "__main__":
    sys.exit(main())
